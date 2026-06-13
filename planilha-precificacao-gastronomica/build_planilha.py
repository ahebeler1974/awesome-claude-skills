# -*- coding: utf-8 -*-
"""
Gera a "Planilha Universal de Precificacao Gastronomica" (.xlsx)
5 abas integradas + aba de configuracao oculta. Formulas, validacao de
dados, formatacao condicional, graficos e paleta profissional.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import DoughnutChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# PALETA DE CORES (moderna, gastronomica, fora do padrao Excel)
# ----------------------------------------------------------------------------
CREAM     = "FAF8F5"  # fundo off-white cremoso
GRAPHITE  = "3A3A38"  # texto / cabecalhos escuros
LIGHTGRAY = "EFEDE8"  # zebra / linhas suaves
MIDGRAY   = "D9D5CD"  # bordas
TERRA     = "C98A6B"  # destaque (terracota / caramelo)
TERRA_DK  = "A9613F"  # destaque escuro
SAGE      = "8A9A7B"  # verde salvia (positivo)
SAGE_LT   = "DCE6D4"  # verde claro (ok)
ALERT     = "C0584E"  # alerta (vermelho suave)
ALERT_LT  = "F4DAD5"  # alerta claro
AMBER_LT  = "F6E9CF"  # atencao
WHITE     = "FFFFFF"

FONT = "Calibri"

# ---- Formatos numericos
BRL  = u'R$ #,##0.00'
PCT  = u'0.0%'
PCT0 = u'0%'
NUM  = u'#,##0.00'
INT  = u'#,##0'

# ---- Estilos base
def f(sz=11, b=False, color=GRAPHITE, italic=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=italic)

def fill(c):
    return PatternFill("solid", fgColor=c)

thin = Side(style="thin", color=MIDGRAY)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NOBORDER = Border()

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",  vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right", vertical="center")
LEFTTOP= Alignment(horizontal="left", vertical="top", wrap_text=True)

LOCKED   = Protection(locked=True)
UNLOCKED = Protection(locked=False)

wb = Workbook()

# ============================================================================
# Helpers
# ============================================================================
def cell(ws, ref, value=None, font=None, bg=None, align=None, border=None,
         numfmt=None, locked=True):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font or f()
    if bg:
        c.fill = fill(bg)
    c.alignment = align or LEFT
    c.border = border if border is not None else NOBORDER
    if numfmt:
        c.number_format = numfmt
    c.protection = LOCKED if locked else UNLOCKED
    return c

def fill_range(ws, rng, bg=None, font=None, align=None, border=None):
    """Aplica estilo a todas as celulas de um range (para fundos de merges)."""
    for row in ws[rng]:
        for c in row:
            if bg:
                c.fill = fill(bg)
            if font:
                c.font = font
            if align:
                c.alignment = align
            if border is not None:
                c.border = border

def banner(ws, rng, title, sub_rng=None, subtitle=None):
    fill_range(ws, rng, bg=GRAPHITE)
    top = rng.split(":")[0]
    cell(ws, top, title, font=f(18, True, WHITE), bg=GRAPHITE, align=LEFT)
    ws.merge_cells(rng)
    if sub_rng and subtitle:
        fill_range(ws, sub_rng, bg=TERRA)
        t2 = sub_rng.split(":")[0]
        cell(ws, t2, subtitle, font=f(10, False, WHITE), bg=TERRA, align=LEFT)
        ws.merge_cells(sub_rng)

def section(ws, rng, text):
    fill_range(ws, rng, bg=TERRA)
    top = rng.split(":")[0]
    cell(ws, top, text, font=f(12, True, WHITE), bg=TERRA, align=LEFT)
    ws.merge_cells(rng)

def add_name(name, attr):
    dn = DefinedName(name, attr_text=attr)
    try:
        wb.defined_names.add(dn)
    except Exception:
        wb.defined_names[name] = dn

def no_grid(ws):
    ws.sheet_view.showGridLines = False

# ============================================================================
# ABA _Config (oculta) - listas e tabela de conversao
# ============================================================================
cfg = wb.active
cfg.title = "_Config"
no_grid(cfg)
cell(cfg, "A1", "Categorias", font=f(10, True))
for i, v in enumerate(["Doce", "Salgado", "Embalagem"], start=2):
    cell(cfg, f"A{i}", v)
cell(cfg, "C1", "Unidades", font=f(10, True))
for i, v in enumerate(["g", "kg", "ml", "L", "un"], start=2):
    cell(cfg, f"C{i}", v)
cell(cfg, "E1", "Tipos de Produto", font=f(10, True))
for i, v in enumerate(["Doce", "Salgado", "Torta"], start=2):
    cell(cfg, f"E{i}", v)
cell(cfg, "G1", "Unidade", font=f(10, True))
cell(cfg, "H1", "Fator p/ Unidade Minima", font=f(10, True))
conv = [("g", 1), ("kg", 1000), ("ml", 1), ("L", 1000), ("un", 1)]
for i, (u, fa) in enumerate(conv, start=2):
    cell(cfg, f"G{i}", u)
    cell(cfg, f"H{i}", fa)

add_name("lista_categorias", "_Config!$A$2:$A$4")
add_name("lista_unidades",   "_Config!$C$2:$C$6")
add_name("lista_tipos",      "_Config!$E$2:$E$4")
add_name("tab_conversao",    "_Config!$G$2:$H$6")

# ============================================================================
# ABA 1 - Banco de Dados de Insumos
# ============================================================================
s1 = wb.create_sheet("1. Insumos")
no_grid(s1)
s1.sheet_properties.tabColor = TERRA
widths1 = {"A": 30, "B": 16, "C": 14, "D": 12, "E": 14, "F": 16, "G": 22}
for col, w in widths1.items():
    s1.column_dimensions[col].width = w

banner(s1, "A1:G1", "BANCO DE DADOS DE INSUMOS",
       "A2:G2", "O motor da planilha — cadastre aqui TODOS os ingredientes e materiais. O custo unitario e calculado sozinho.")
s1.row_dimensions[1].height = 30

headers1 = ["Ingrediente / Insumo", "Categoria", "Qtd da Embalagem de Compra",
            "Unidade de Medida", "Preço Pago (R$)", "Fator de Rendimento (%)",
            "Custo por Unidade Mínima (R$)"]
for i, h in enumerate(headers1):
    col = get_column_letter(i + 1)
    cell(s1, f"{col}3", h, font=f(10, True, WHITE), bg=GRAPHITE, align=CENTER, border=BORDER)
s1.row_dimensions[3].height = 38

insumos = [
    ("Leite Condensado",      "Doce",      395, "g",  6.50, 1.00),
    ("Creme de Leite",        "Doce",      200, "g",  3.80, 1.00),
    ("Chocolate em Pó 50%",   "Doce",     1000, "g", 28.00, 1.00),
    ("Manteiga",              "Doce",      500, "g", 22.00, 1.00),
    ("Granulado Belga",       "Doce",      500, "g", 35.00, 1.00),
    ("Açúcar Refinado",       "Doce",     1000, "g",  4.50, 1.00),
    ("Leite Integral",        "Doce",     1000, "ml", 5.00, 1.00),
    ("Farinha de Trigo",      "Salgado",  1000, "g",  4.80, 1.00),
    ("Frango (peito)",        "Salgado",  1000, "g", 18.00, 0.85),
    ("Queijo Muçarela",       "Salgado",  1000, "g", 32.00, 0.95),
    ("Óleo de Soja",          "Salgado",   900, "ml", 7.20, 1.00),
    ("Ovos",                  "Salgado",    12, "un", 9.00, 1.00),
    ("Farinha de Rosca",      "Salgado",   500, "g",  6.00, 1.00),
    ("Margarina",             "Salgado",   500, "g",  8.00, 1.00),
    ("Fermento Químico",      "Doce",      100, "g",  5.50, 1.00),
    ("Forminha Nº4",          "Embalagem", 100, "un", 4.50, 1.00),
    ("Caixa Kraft 6 doces",   "Embalagem",  50, "un",45.00, 1.00),
    ("Embalagem Torta",       "Embalagem",  25, "un",30.00, 1.00),
]
FIRST = 4
LAST = 203
for r in range(FIRST, LAST + 1):
    idx = r - FIRST
    zebra = LIGHTGRAY if (idx % 2 == 1) else WHITE
    data = insumos[idx] if idx < len(insumos) else None
    # A..F entradas (desbloqueadas)
    cell(s1, f"A{r}", data[0] if data else None, bg=zebra, align=LEFT, border=BORDER, locked=False)
    cell(s1, f"B{r}", data[1] if data else None, bg=zebra, align=CENTER, border=BORDER, locked=False)
    cell(s1, f"C{r}", data[2] if data else None, bg=zebra, align=CENTER, border=BORDER, numfmt=NUM, locked=False)
    cell(s1, f"D{r}", data[3] if data else None, bg=zebra, align=CENTER, border=BORDER, locked=False)
    cell(s1, f"E{r}", data[4] if data else None, bg=zebra, align=RIGHT, border=BORDER, numfmt=BRL, locked=False)
    cell(s1, f"F{r}", data[5] if data else None, bg=zebra, align=CENTER, border=BORDER, numfmt=PCT0, locked=False)
    # G automatica (bloqueada, fundo distinto)
    formula = (f"=IFERROR(E{r}/(C{r}*VLOOKUP(D{r},tab_conversao,2,0)*"
               f"IF(OR(F{r}=\"\",F{r}=0),1,F{r})),0)")
    cell(s1, f"G{r}", formula, bg="F1EBE3", align=RIGHT, border=BORDER, numfmt=BRL,
         font=f(11, False, TERRA_DK, italic=True), locked=True)

s1.freeze_panes = "A4"

# Validacao de dados
dv_cat = DataValidation(type="list", formula1="lista_categorias", allow_blank=True)
dv_uni = DataValidation(type="list", formula1="lista_unidades", allow_blank=True)
s1.add_data_validation(dv_cat); dv_cat.add(f"B{FIRST}:B{LAST}")
s1.add_data_validation(dv_uni); dv_uni.add(f"D{FIRST}:D{LAST}")

add_name("tab_insumos",  "'1. Insumos'!$A$4:$G$203")
add_name("lista_insumos","'1. Insumos'!$A$4:$A$203")
s1.protection.sheet = True

# ============================================================================
# ABA 2 - Ficha Tecnica Universal
# ============================================================================
s2 = wb.create_sheet("2. Ficha Técnica")
no_grid(s2)
s2.sheet_properties.tabColor = SAGE
for col, w in {"A": 32, "B": 18, "C": 14, "D": 20}.items():
    s2.column_dimensions[col].width = w

banner(s2, "A1:D1", "FICHA TÉCNICA UNIVERSAL",
       "A2:D2", "Monte a receita escolhendo os insumos no menu suspenso. O custo de cada item é puxado automaticamente da Aba 1.")
s2.row_dimensions[1].height = 30

cell(s2, "A3", "Nome da Receita:", font=f(11, True), bg=LIGHTGRAY, align=RIGHT, border=BORDER)
fill_range(s2, "B3:D3", bg=WHITE, border=BORDER)
cell(s2, "B3", "Brigadeiro Gourmet", font=f(11, True, TERRA_DK), align=LEFT, border=BORDER, locked=False)
ws_merge = s2.merge_cells("B3:D3")
cell(s2, "A4", "Tipo de Produto:", font=f(11, True), bg=LIGHTGRAY, align=RIGHT, border=BORDER)
cell(s2, "B4", "Doce", font=f(11, True, TERRA_DK), bg=WHITE, align=LEFT, border=BORDER, locked=False)
cell(s2, "C4", "← Doce / Salgado / Torta", font=f(9, False, GRAPHITE, italic=True), align=LEFT)
s2.merge_cells("C4:D4")

headers2 = ["Item da Receita", "Quantidade Utilizada", "Unidade", "Custo Calculado (R$)"]
HR = 6
for i, h in enumerate(headers2):
    col = get_column_letter(i + 1)
    cell(s2, f"{col}{HR}", h, font=f(10, True, WHITE), bg=GRAPHITE, align=CENTER, border=BORDER)
s2.row_dimensions[HR].height = 30

ficha = [
    ("Leite Condensado",   395, "g"),
    ("Creme de Leite",     100, "g"),
    ("Chocolate em Pó 50%", 60, "g"),
    ("Manteiga",            15, "g"),
    ("Granulado Belga",    150, "g"),
]
DSTART = HR + 1     # 7
DEND = DSTART + 29  # 36
for r in range(DSTART, DEND + 1):
    idx = r - DSTART
    zebra = LIGHTGRAY if (idx % 2 == 1) else WHITE
    d = ficha[idx] if idx < len(ficha) else None
    cell(s2, f"A{r}", d[0] if d else None, bg=zebra, align=LEFT, border=BORDER, locked=False)
    cell(s2, f"B{r}", d[1] if d else None, bg=zebra, align=CENTER, border=BORDER, numfmt=NUM, locked=False)
    cell(s2, f"C{r}", d[2] if d else None, bg=zebra, align=CENTER, border=BORDER, locked=False)
    formula = f"=IFERROR(VLOOKUP(A{r},tab_insumos,7,0)*B{r}*VLOOKUP(C{r},tab_conversao,2,0),0)"
    cell(s2, f"D{r}", formula, bg="EDF0EA", align=RIGHT, border=BORDER, numfmt=BRL,
         font=f(11, False, SAGE, italic=True), locked=True)

TR = DEND + 1  # 37
fill_range(s2, f"A{TR}:C{TR}", bg=GRAPHITE)
cell(s2, f"A{TR}", "CUSTO TOTAL DE INGREDIENTES (por receita)", font=f(11, True, WHITE), bg=GRAPHITE, align=RIGHT, border=BORDER)
s2.merge_cells(f"A{TR}:C{TR}")
cell(s2, f"D{TR}", f"=SUM(D{DSTART}:D{DEND})", font=f(12, True, WHITE), bg=TERRA_DK, align=RIGHT, border=BORDER, numfmt=BRL)

s2.freeze_panes = f"A{DSTART}"

dv_ins = DataValidation(type="list", formula1="lista_insumos", allow_blank=True)
dv_uni2 = DataValidation(type="list", formula1="lista_unidades", allow_blank=True)
dv_tipo = DataValidation(type="list", formula1="lista_tipos", allow_blank=True)
s2.add_data_validation(dv_ins);  dv_ins.add(f"A{DSTART}:A{DEND}")
s2.add_data_validation(dv_uni2); dv_uni2.add(f"C{DSTART}:C{DEND}")
s2.add_data_validation(dv_tipo); dv_tipo.add("B4")

add_name("custo_ingredientes", "'2. Ficha Técnica'!$D$37")
add_name("receita_nome",       "'2. Ficha Técnica'!$B$3")
add_name("receita_tipo",       "'2. Ficha Técnica'!$B$4")
s2.protection.sheet = True

# ============================================================================
# ABA 3 - Custos & Mao de Obra
# ============================================================================
s3 = wb.create_sheet("3. Custos & Mão de Obra")
no_grid(s3)
s3.sheet_properties.tabColor = TERRA_DK
for col, w in {"A": 4, "B": 42, "C": 18, "D": 40}.items():
    s3.column_dimensions[col].width = w

banner(s3, "A1:D1", "MÃO DE OBRA E CUSTOS INVISÍVEIS",
       "A2:D2", "Configure uma vez. A planilha transforma tudo em custo por minuto e aplica em cada receita automaticamente.")
s3.row_dimensions[1].height = 30

def row3(r, label, value=None, numfmt=None, auto=False, note=None, name=None):
    bg = "F1EBE3" if auto else WHITE
    cell(s3, f"B{r}", label, font=f(11, True if auto else False), bg=LIGHTGRAY if not auto else "F1EBE3",
         align=RIGHT, border=BORDER)
    cell(s3, f"C{r}", value, bg=bg, align=RIGHT, border=BORDER, numfmt=numfmt,
         font=f(11, auto, TERRA_DK if auto else GRAPHITE, italic=auto), locked=not (value is not None and not auto) and False if not auto else True)
    # entradas: desbloqueadas; automaticas: bloqueadas
    s3[f"C{r}"].protection = LOCKED if auto else UNLOCKED
    if note:
        cell(s3, f"D{r}", note, font=f(9, False, GRAPHITE, italic=True), align=LEFT)

section(s3, "A3:D3", "BLOCO 1 — VALOR DO SEU TEMPO (MÃO DE OBRA)")
row3(4,  "Pró-labore desejado (R$/mês)", 2500, BRL, note="Quanto você quer receber por mês.")
row3(5,  "Horas trabalhadas por dia", 6, NUM, note="Média de horas dedicadas à produção.")
row3(6,  "Dias trabalhados por mês", 22, NUM)
row3(7,  "Horas por mês", "=C5*C6", NUM, auto=True)
row3(8,  "Minutos por mês", "=C7*60", INT, auto=True, note="Base para ratear todos os custos por minuto.")
row3(9,  "Valor da Hora (R$)", "=IFERROR(C4/C7,0)", BRL, auto=True)
row3(10, "Valor do Minuto de Mão de Obra (R$)", "=IFERROR(C9/60,0)", BRL, auto=True)

section(s3, "A12:D12", "BLOCO 2 — ENERGIA & GÁS (EQUIPAMENTOS)")
row3(13, "Gás (R$/mês)", 90, BRL)
row3(14, "Energia Elétrica (R$/mês)", 180, BRL)
row3(15, "Minutos de forno/fogão por mês (estimativa)", 2400, INT, note="Quanto tempo seus equipamentos ficam ligados no mês.")
row3(16, "Custo por Minuto de Forno/Fogão (R$)", "=IFERROR((C13+C14)/C15,0)", BRL, auto=True)

section(s3, "A18:D18", "BLOCO 3 — CUSTOS FIXOS DE ESTRUTURA (R$/mês)")
row3(19, "Água", 60, BRL)
row3(20, "MEI / Impostos fixos", 76, BRL)
row3(21, "Aluguel / Espaço", 0, BRL)
row3(22, "Internet / Telefone", 50, BRL)
row3(23, "Outros (contador, marketing...)", 40, BRL)
row3(24, "Total de Custos Fixos de Estrutura", "=SUM(C19:C23)", BRL, auto=True)
row3(25, "Custo Fixo de Estrutura por Minuto (R$)", "=IFERROR(C24/C8,0)", BRL, auto=True,
     note="Rateado pelos minutos trabalhados no mês.")

add_name("valor_min_mo",       "'3. Custos & Mão de Obra'!$C$10")
add_name("custo_min_forno",    "'3. Custos & Mão de Obra'!$C$16")
add_name("custo_min_estrutura","'3. Custos & Mão de Obra'!$C$25")
s3.protection.sheet = True

# ============================================================================
# ABA 4 - Simulador de Margem e Preco Final (DASHBOARD)
# ============================================================================
s4 = wb.create_sheet("4. Precificação")
no_grid(s4)
s4.sheet_properties.tabColor = GRAPHITE
for col, w in {"A": 40, "B": 18, "C": 20, "D": 22, "E": 20, "F": 16, "G": 4}.items():
    s4.column_dimensions[col].width = w

banner(s4, "A1:G1", "PRECIFICAÇÃO — PAINEL DE VENDAS",
       "A2:G2", "Sua tela principal. Preencha apenas os campos claros. Tudo se calcula sozinho.")
s4.row_dimensions[1].height = 30

def lbl(r, text, col="A"):
    cell(s4, f"{col}{r}", text, font=f(11, False), bg=LIGHTGRAY, align=RIGHT, border=BORDER)

def inp(r, value, numfmt=None, col="C"):
    cell(s4, f"{col}{r}", value, bg=WHITE, align=RIGHT, border=BORDER, numfmt=numfmt,
         font=f(11, True, TERRA_DK), locked=False)

def auto(r, formula, numfmt=BRL, col="C", big=False):
    cell(s4, f"{col}{r}", formula, bg="F1EBE3", align=RIGHT, border=BORDER, numfmt=numfmt,
         font=f(13 if big else 11, True, GRAPHITE, italic=not big), locked=True)

section(s4, "A3:G3", "DADOS DA RECEITA")
lbl(4, "Receita")
cell(s4, "B4", "=receita_nome", bg=WHITE, align=LEFT, border=BORDER, font=f(11, True, TERRA_DK))
s4.merge_cells("B4:D4")
lbl(5, "Tipo de Produto")
cell(s4, "B5", "=receita_tipo", bg=WHITE, align=LEFT, border=BORDER, font=f(11, True, TERRA_DK))
lbl(6, "Rendimento da Receita (unidades)")
inp(6, 30, INT)
cell(s4, "D6", "← Ex.: 50 doces, 100 salgados ou 2 tortas", font=f(9, False, GRAPHITE, italic=True), align=LEFT)
lbl(7, "Tempo de Produção / Mão de Obra (min)")
inp(7, 40, INT)
lbl(8, "Tempo de Forno / Fogão (min)")
inp(8, 15, INT)
cell(s4, "D8", "← 0 para doces sem cocção", font=f(9, False, GRAPHITE, italic=True), align=LEFT)

section(s4, "A10:D10", "EMBALAGEM (por unidade do produto)")
for i, h in enumerate(["Item de Embalagem", "Qtd por unidade", "Custo (R$)"]):
    col = get_column_letter(i + 1)
    cell(s4, f"{col}11", h, font=f(10, True, WHITE), bg=GRAPHITE, align=CENTER, border=BORDER)
emb_rows = [("Forminha Nº4", 1), ("Caixa Kraft 6 doces", 0.1667), (None, None)]
for k, (item, qty) in enumerate(emb_rows):
    r = 12 + k
    cell(s4, f"A{r}", item, bg=WHITE, align=LEFT, border=BORDER, locked=False)
    cell(s4, f"B{r}", qty, bg=WHITE, align=CENTER, border=BORDER, numfmt=NUM, locked=False)
    cell(s4, f"C{r}", f"=IFERROR(VLOOKUP(A{r},tab_insumos,7,0)*B{r},0)", bg="F1EBE3",
         align=RIGHT, border=BORDER, numfmt=BRL, font=f(11, False, TERRA_DK, italic=True))
cell(s4, "A15", "Custo de Embalagem por Unidade (R$)", font=f(11, True, WHITE), bg=TERRA_DK, align=RIGHT, border=BORDER)
s4.merge_cells("A15:B15")
cell(s4, "C15", "=SUM(C12:C14)", bg=TERRA_DK, align=RIGHT, border=BORDER, numfmt=BRL, font=f(11, True, WHITE))

dv_emb = DataValidation(type="list", formula1="lista_insumos", allow_blank=True)
s4.add_data_validation(dv_emb); dv_emb.add("A12:A14")

section(s4, "A17:G17", "CUSTOS CONSOLIDADOS (por receita)")
lbl(18, "Ingredientes (Ficha Técnica)");        auto(18, "=custo_ingredientes")
lbl(19, "Mão de Obra");                          auto(19, "=tempo_producao*valor_min_mo")
lbl(20, "Energia / Gás (Forno/Fogão)");          auto(20, "=tempo_forno*custo_min_forno")
lbl(21, "Custos Fixos (Estrutura)");             auto(21, "=tempo_producao*custo_min_estrutura")
lbl(22, "Embalagem (receita inteira)");          auto(22, "=emb_unit*rendimento")
cell(s4, "A23", "CUSTO TOTAL DA RECEITA", font=f(11, True, WHITE), bg=GRAPHITE, align=RIGHT, border=BORDER)
cell(s4, "C23", "=SUM(C18:C22)", bg=GRAPHITE, align=RIGHT, border=BORDER, numfmt=BRL, font=f(11, True, WHITE))
cell(s4, "A24", "CUSTO TOTAL POR UNIDADE", font=f(12, True, WHITE), bg=TERRA_DK, align=RIGHT, border=BORDER)
cell(s4, "C24", "=IFERROR(C23/rendimento,0)", bg=TERRA_DK, align=RIGHT, border=BORDER, numfmt=BRL, font=f(13, True, WHITE))

section(s4, "A26:G26", "PRECIFICAÇÃO")
lbl(27, "Margem de Lucro Desejada (%)");   inp(27, 0.50, PCT)
lbl(28, "Margem Mínima Saudável (%)");     inp(28, 0.20, PCT)
cell(s4, "D28", "← Abaixo disto a planilha te alerta em vermelho", font=f(9, False, GRAPHITE, italic=True), align=LEFT)
lbl(29, "Preço Mínimo de Venda / un (à vista, sem taxa)")
auto(29, "=IFERROR(custo_unit/(1-margem),0)", BRL, big=True)
lbl(30, "Lucro por unidade (à vista)")
auto(30, "=C29-custo_unit", BRL)

section(s4, "A32:F32", "CENÁRIOS DE CANAIS DE VENDA")
ch = ["Canal", "Taxa do Canal (%)", "Preço Sugerido (R$)", "Preço Final Arredondado (R$)",
      "Lucro Líquido / un (R$)", "Margem Real (%)"]
for i, h in enumerate(ch):
    col = get_column_letter(i + 1)
    cell(s4, f"{col}33", h, font=f(10, True, WHITE), bg=GRAPHITE, align=CENTER, border=BORDER)
s4.row_dimensions[33].height = 30
canais = [("Balcão (Dinheiro / PIX)", 0.0), ("Cartão (Maquininha)", 0.035), ("Aplicativo (iFood)", 0.27)]
for k, (nome, taxa) in enumerate(canais):
    r = 34 + k
    cell(s4, f"A{r}", nome, bg=WHITE, align=LEFT, border=BORDER, font=f(11, True))
    cell(s4, f"B{r}", taxa, bg=WHITE, align=RIGHT, border=BORDER, numfmt=PCT, locked=False, font=f(11, True, TERRA_DK))
    cell(s4, f"C{r}", f"=IFERROR(custo_unit/(1-margem-B{r}),0)", bg="F1EBE3", align=RIGHT, border=BORDER, numfmt=BRL, font=f(11, False, GRAPHITE, italic=True))
    cell(s4, f"D{r}", f"=IFERROR(CEILING(C{r},0.5),0)", bg="F1EBE3", align=RIGHT, border=BORDER, numfmt=BRL, font=f(12, True, GRAPHITE))
    cell(s4, f"E{r}", f"=D{r}*(1-B{r})-custo_unit", bg="F1EBE3", align=RIGHT, border=BORDER, numfmt=BRL, font=f(11, False, GRAPHITE, italic=True))
    cell(s4, f"F{r}", f"=IFERROR((D{r}*(1-B{r})-custo_unit)/D{r},0)", bg="F1EBE3", align=RIGHT, border=BORDER, numfmt=PCT, font=f(11, True, GRAPHITE))

# Nomes do dashboard
add_name("rendimento",     "'4. Precificação'!$C$6")
add_name("tempo_producao", "'4. Precificação'!$C$7")
add_name("tempo_forno",    "'4. Precificação'!$C$8")
add_name("emb_unit",       "'4. Precificação'!$C$15")
add_name("custo_unit",     "'4. Precificação'!$C$24")
add_name("margem",         "'4. Precificação'!$C$27")
add_name("margem_minima",  "'4. Precificação'!$C$28")
add_name("preco_min",      "'4. Precificação'!$C$29")
add_name("preco_ref",      "'4. Precificação'!$D$35")  # Cartao - referencia do grafico
add_name("taxa_ref",       "'4. Precificação'!$B$35")

# Formatacao condicional
alert_fill = fill(ALERT_LT); alert_font = f(11, True, ALERT)
ok_fill = fill(SAGE_LT);     ok_font = f(11, True, "4F663F")
# Margem desejada abaixo da minima
s4.conditional_formatting.add("C27", FormulaRule(formula=["$C$27<$C$28"], fill=alert_fill, font=alert_font, stopIfTrue=False))
# Margem real por canal
s4.conditional_formatting.add("F34:F36", FormulaRule(formula=["F34<$C$28"], fill=alert_fill, font=alert_font))
s4.conditional_formatting.add("F34:F36", FormulaRule(formula=["F34>=$C$28"], fill=ok_fill, font=ok_font))
# Lucro negativo
s4.conditional_formatting.add("E34:E36", FormulaRule(formula=["E34<0"], fill=alert_fill, font=alert_font))
s4.conditional_formatting.add("C30", FormulaRule(formula=["$C$30<0"], fill=alert_fill, font=alert_font))

s4.protection.sheet = True

# ============================================================================
# ABA 5 - Painel Visual (composicao + grafico + guia de design)
# ============================================================================
s5 = wb.create_sheet("5. Painel Visual")
no_grid(s5)
s5.sheet_properties.tabColor = TERRA
for col, w in {"A": 34, "B": 16, "C": 12, "D": 3, "E": 30, "F": 30, "G": 22}.items():
    s5.column_dimensions[col].width = w

banner(s5, "A1:G1", "PAINEL VISUAL — COMPOSIÇÃO DO SEU PREÇO",
       "A2:G2", "Mostra exatamente para onde vai cada centavo do preço de venda (base: canal Cartão).")
s5.row_dimensions[1].height = 30

section(s5, "A4:C4", "DE QUE É FEITO O SEU PREÇO (por unidade)")
for i, h in enumerate(["Componente", "Valor / un (R$)", "% do Preço"]):
    col = get_column_letter(i + 1)
    cell(s5, f"{col}5", h, font=f(10, True, WHITE), bg=GRAPHITE, align=CENTER, border=BORDER)

comps = [
    ("Ingredientes",                "=IFERROR(custo_ingredientes/rendimento,0)", "F1EBE3"),
    ("Embalagem",                   "=emb_unit", "F1EBE3"),
    ("Mão de Obra",                 "=IFERROR(tempo_producao*valor_min_mo/rendimento,0)", "F1EBE3"),
    ("Custos Fixos (Energia+Estr.)","=IFERROR((tempo_forno*custo_min_forno+tempo_producao*custo_min_estrutura)/rendimento,0)", "F1EBE3"),
    ("Taxas do Canal (Cartão)",     "=preco_ref*taxa_ref", "F1EBE3"),
    ("Lucro Real",                  "=preco_ref-SUM(B6:B10)", "DCE6D4"),
]
for k, (nome, formula, bg) in enumerate(comps):
    r = 6 + k
    cell(s5, f"A{r}", nome, bg=WHITE, align=LEFT, border=BORDER, font=f(11, k == 5, GRAPHITE))
    cell(s5, f"B{r}", formula, bg=bg, align=RIGHT, border=BORDER, numfmt=BRL, font=f(11, k == 5, GRAPHITE))
    cell(s5, f"C{r}", f"=IFERROR(B{r}/preco_ref,0)", bg=bg, align=RIGHT, border=BORDER, numfmt=PCT, font=f(11, k == 5, GRAPHITE))
cell(s5, "A12", "PREÇO FINAL (Cartão)", font=f(12, True, WHITE), bg=TERRA_DK, align=RIGHT, border=BORDER)
cell(s5, "B12", "=preco_ref", bg=TERRA_DK, align=RIGHT, border=BORDER, numfmt=BRL, font=f(12, True, WHITE))
cell(s5, "C12", "=IFERROR(B12/preco_ref,0)", bg=TERRA_DK, align=RIGHT, border=BORDER, numfmt=PCT0, font=f(12, True, WHITE))

# Grafico de rosca (doughnut)
dough = DoughnutChart()
dough.title = "Composição do Preço (R$/un)"
labels = Reference(s5, min_col=1, min_row=6, max_row=11)
data = Reference(s5, min_col=2, min_row=5, max_row=11)
dough.add_data(data, titles_from_data=True)
dough.set_categories(labels)
dough.height = 8.5
dough.width = 13
dough.holeSize = 55
dough.dataLabels = DataLabelList(); dough.dataLabels.showPercent = True
s5.add_chart(dough, "E4")

# Grafico de barras
bar = BarChart()
bar.type = "bar"
bar.title = "Quanto pesa cada componente"
bar.add_data(data, titles_from_data=True)
bar.set_categories(labels)
bar.height = 8.5
bar.width = 13
bar.legend = None
s5.add_chart(bar, "E20")

# Guia de design (texto)
GR = 15
section(s5, f"A{GR}:C{GR}", "GUIA DE DESIGN DESTA PLANILHA")
guia = [
    ("Paleta de cores", ""),
    ("  Fundo (off-white cremoso)", "#FAF8F5"),
    ("  Cabeçalhos (grafite)", "#3A3A38"),
    ("  Destaque (terracota/caramelo)", "#C98A6B"),
    ("  Linhas/zebra (cinza claro)", "#EFEDE8"),
    ("  Positivo (verde sálvia)", "#8A9A7B"),
    ("  Alerta (vermelho suave)", "#C0584E"),
    ("Formatação condicional", ""),
    ("  Margem < mínima saudável", "Célula fica VERMELHA"),
    ("  Margem >= mínima saudável", "Célula fica VERDE"),
    ("  Lucro por unidade negativo", "Alerta VERMELHO"),
    ("Gráficos", ""),
    ("  Rosca: % de cada componente", "no preço final"),
    ("  Barras: peso em R$ de cada parte", "do preço"),
]
for k, (a, b) in enumerate(guia):
    r = GR + 1 + k
    is_head = b == "" and not a.startswith("  ")
    cell(s5, f"A{r}", a, font=f(11, is_head, TERRA_DK if is_head else GRAPHITE),
         bg=LIGHTGRAY if is_head else WHITE, align=LEFT, border=BORDER)
    cell(s5, f"B{r}", b, font=f(10, False, GRAPHITE), bg=LIGHTGRAY if is_head else WHITE, align=LEFT, border=BORDER)
    s5.merge_cells(f"B{r}:C{r}")

# ============================================================================
# Ocultar config e salvar
# ============================================================================
cfg.sheet_state = "hidden"
wb.active = s4  # abrir no dashboard

OUT = "/home/user/awesome-claude-skills/Planilha_Universal_Precificacao_Gastronomica.xlsx"
wb.save(OUT)
print("OK:", OUT)
print("Abas:", [ws.title for ws in wb.worksheets])
print("Nomes definidos:", sorted(list(wb.defined_names)) if hasattr(wb.defined_names, "__iter__") else "n/a")
